"""Independently read the browser-generated XLSX with openpyxl and ZIP/XML parsers."""
from pathlib import Path
import json
import sys
from datetime import datetime
import zipfile
import xml.etree.ElementTree as ET
import openpyxl
root = Path(__file__).resolve().parents[1]
book_path = Path(sys.argv[1]) if len(sys.argv)>1 else root / 'verification' / 'example-snapshot.xlsx'
expected_path = Path(sys.argv[2]) if len(sys.argv)>2 else root / 'verification' / 'expected-export.json'
expected = json.loads(expected_path.read_text())
with zipfile.ZipFile(book_path) as archive:
    assert archive.testzip() is None
    for name in archive.namelist():
        if name.endswith(('.xml', '.rels')):
            ET.fromstring(archive.read(name))
book = openpyxl.load_workbook(book_path, data_only=False)
assert book.sheetnames == [s['name'] for s in expected]
count = 0
for sheet in expected:
    actual = book[sheet['name']]
    assert actual.freeze_panes == 'A2'
    for ri, row in enumerate(sheet['rows'], 1):
        for ci, value in enumerate(row, 1):
            cell = actual.cell(ri, ci)
            assert cell.data_type != 'f', 'Snapshot must contain no formulas'
            if row[0] == 'Exported UTC' and ci == 2:
                # Browser export and pre-export snapshot are separate calls.
                delta = datetime.fromisoformat(cell.value.replace('Z', '+00:00')) - datetime.fromisoformat(value.replace('Z', '+00:00'))
                assert abs(delta.total_seconds()) < 5
            elif value == '':
                assert cell.value in ('', None)
            elif isinstance(value, (int, float)):
                assert isinstance(cell.value, (int, float))
                assert abs(cell.value-value) < 0.000001, (sheet['name'], ri, ci)
            else:
                assert cell.value == value, (sheet['name'], ri, ci, cell.value, value)
            if ri > 1 and isinstance(value, (int, float)) and ci-1 in sheet.get('currencyColumns', []):
                assert '$' in cell.number_format
            count += 1
print(f'PASS: {len(book.sheetnames)} sheets, {count} cells, valid ZIP/XML, typed values, currency formatting, no formulas.')
